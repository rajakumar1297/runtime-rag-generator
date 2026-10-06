"""
Document ingestion module.
Handles parsing of various document formats (PDF, DOCX, TXT, CSV, MD)
and splitting them into chunks for embedding.
"""

import os
import csv
import logging
from pathlib import Path
from typing import List

from langchain.schema import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.config import settings

logger = logging.getLogger(__name__)

# Supported file extensions and their MIME types
SUPPORTED_EXTENSIONS = {
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".csv": "text/csv",
}


def parse_pdf(file_path: str) -> str:
    """Extract text from a PDF file."""
    from pypdf import PdfReader

    reader = PdfReader(file_path)
    text_parts = []
    for i, page in enumerate(reader.pages):
        page_text = page.extract_text()
        if page_text:
            text_parts.append(f"[Page {i + 1}]\n{page_text}")
    return "\n\n".join(text_parts)


def parse_docx(file_path: str) -> str:
    """Extract text from a DOCX file."""
    from docx import Document as DocxDocument

    doc = DocxDocument(file_path)
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
    return "\n\n".join(paragraphs)


def parse_txt(file_path: str) -> str:
    """Read a plain text or markdown file."""
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def parse_csv(file_path: str) -> str:
    """Convert CSV to readable text format."""
    rows = []
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f)
        headers = next(reader, None)
        if headers:
            for row in reader:
                row_text = ", ".join(
                    f"{h}: {v}" for h, v in zip(headers, row) if v.strip()
                )
                if row_text:
                    rows.append(row_text)
    return "\n".join(rows)


# Map extensions to their parser functions
PARSERS = {
    ".pdf": parse_pdf,
    ".txt": parse_txt,
    ".md": parse_txt,
    ".docx": parse_docx,
    ".csv": parse_csv,
}


def parse_document(file_path: str) -> str:
    """
    Parse a document file and return its text content.

    Args:
        file_path: Path to the document file.

    Returns:
        Extracted text content.

    Raises:
        ValueError: If the file format is not supported.
    """
    ext = Path(file_path).suffix.lower()
    parser = PARSERS.get(ext)
    if not parser:
        raise ValueError(
            f"Unsupported file format: {ext}. "
            f"Supported formats: {', '.join(SUPPORTED_EXTENSIONS.keys())}"
        )
    logger.info(f"Parsing document: {file_path} (format: {ext})")
    return parser(file_path)


def create_documents_from_text(
    text: str, filename: str, collection_name: str
) -> List[Document]:
    """
    Split text into chunks and create LangChain Document objects.

    Args:
        text: The raw text content.
        filename: Original filename for metadata.
        collection_name: The collection this document belongs to.

    Returns:
        List of LangChain Document objects with metadata.
    """
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        length_function=len,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    chunks = splitter.split_text(text)
    documents = []
    for i, chunk in enumerate(chunks):
        doc = Document(
            page_content=chunk,
            metadata={
                "source": filename,
                "collection": collection_name,
                "chunk_index": i,
                "total_chunks": len(chunks),
            },
        )
        documents.append(doc)

    logger.info(f"Created {len(documents)} chunks from {filename}")
    return documents


def ingest_file(
    file_path: str, filename: str, collection_name: str
) -> List[Document]:
    """
    Full ingestion pipeline: parse a file and split it into document chunks.

    Args:
        file_path: Path to the uploaded file.
        filename: Original filename.
        collection_name: Target collection name.

    Returns:
        List of Document chunks ready for embedding.
    """
    text = parse_document(file_path)
    if not text.strip():
        raise ValueError(f"No text content could be extracted from {filename}")
    return create_documents_from_text(text, filename, collection_name)
